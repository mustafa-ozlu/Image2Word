import os
from PIL import Image
import pytesseract
from docx import Document
import sys

# Çalışma dizinini belirle
def get_resource_path(relative_path):
    """Kaynak dosyalarının doğru şekilde bulunmasını sağlar."""
    if hasattr(sys, "_MEIPASS"):
        return os.path.join(sys._MEIPASS, relative_path)
    return os.path.join(os.path.abspath("."), relative_path)

# Tesseract'ın yolu ve tessdata dizini
pytesseract.pytesseract.tesseract_cmd = get_resource_path("tesseract.exe")
os.environ["TESSDATA_PREFIX"] = get_resource_path("tessdata")

# OCR işlemi
def main():
    image_folder = r"C:\Resimler"
    output_word_file = r"C:\Resimler\OUTPUT.docx"
    doc = Document()

    for filename in os.listdir(image_folder):
        if filename.lower().endswith(('.png', '.jpg', '.jpeg', '.tiff', '.bmp', '.gif')):
            image_path = os.path.join(image_folder, filename)
            print(f"İşleniyor: {image_path}")
            image = Image.open(image_path)
            text = pytesseract.image_to_string(image, lang="tur")
            doc.add_heading(filename, level=1)
            doc.add_paragraph(text)

    doc.save(output_word_file)
    print(f"Tüm işlemler tamamlandı! Çıktı: {output_word_file}")

if __name__ == "__main__":
    main()
